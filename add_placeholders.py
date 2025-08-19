import json
from docx import Document
import re

def extract_words_from_docs():
    """Extract words from Word documents to see which units need placeholders"""
    all_words = {}
    
    for book_num in range(3, 7):
        file_path = f"Book {book_num}.docx"
        print(f"\n=== Analyzing {file_path} ===")
        
        try:
            doc = Document(file_path)
            units = {}
            current_unit = None
            current_words = []
            
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                
                # Check if it's a unit header
                if text.startswith('Unit'):
                    # Save previous unit
                    if current_unit and current_words:
                        units[current_unit] = current_words
                        current_words = []
                    
                    # Extract unit number
                    unit_match = re.search(r'Unit\s+(\d+)', text)
                    if unit_match:
                        current_unit = unit_match.group(1)
                        print(f"  Found Unit {current_unit}")
                    continue
                
                # Check if it's a vocabulary entry
                if current_unit and '–' in text and '(' in text and ')' in text:
                    try:
                        # Split by the en dash
                        parts = text.split('–', 1)
                        if len(parts) == 2:
                            left_part = parts[0].strip()
                            meaning = parts[1].strip()
                            
                            # Extract word and part of speech
                            word_pos_match = re.match(r'^([A-Za-z]+)\s*\(([^)]+)\)', left_part)
                            if word_pos_match:
                                word = word_pos_match.group(1).lower().strip()
                                pos = word_pos_match.group(2).strip().lower()
                                
                                # Clean up part of speech
                                if pos in ['adj', 'adjective']:
                                    pos = "adjective"
                                elif pos in ['v', 'verb']:
                                    pos = "verb"
                                elif pos in ['n', 'noun']:
                                    pos = "noun"
                                elif pos in ['adv', 'adverb']:
                                    pos = "adverb"
                                elif '/' in pos:
                                    pos = pos.split('/')[0].strip()
                                
                                # Skip if word is too short
                                if len(word) < 2:
                                    continue
                                
                                current_words.append({
                                    "word": word,
                                    "partOfSpeech": pos,
                                    "meaning": meaning,
                                    "example": f"The {word} is important for learning."
                                })
                    except:
                        continue
            
            # Add the last unit
            if current_unit and current_words:
                units[current_unit] = current_words
            
            all_words[str(book_num)] = {
                "title": f"Book {book_num} - Academic Vocabulary",
                "units": units
            }
            
            # Print unit statistics
            for unit_num in range(1, 31):
                unit_str = str(unit_num)
                word_count = len(units.get(unit_str, []))
                print(f"    Unit {unit_num}: {word_count} words")
                if word_count < 20:
                    print(f"      ⚠️  Needs {20 - word_count} placeholders")
            
        except Exception as e:
            print(f"Error reading {file_path}: {e}")
            all_words[str(book_num)] = {
                "title": f"Book {book_num} - Academic Vocabulary",
                "units": {}
            }
    
    return all_words

def add_placeholders_to_units():
    """Add placeholder flashcards to units that don't have 20 words"""
    print("Adding placeholders to incomplete units...")
    
    # Get words from documents
    doc_words = extract_words_from_docs()
    
    # Read current app data
    try:
        with open('bookData.js', 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Extract the bookData object
        start = content.find('const bookData = ') + len('const bookData = ')
        end = content.find(';\n\n// Export', start)
        
        if start == -1 or end == -1:
            print("Could not find bookData in file")
            return
        
        # Parse existing data
        app_data = json.loads(content[start:end])
        
        # Process each book
        for book_num in range(3, 7):
            book_str = str(book_num)
            doc_book = doc_words.get(book_str, {"units": {}})
            app_book = app_data.get(book_str, {"units": {}})
            
            print(f"\nProcessing Book {book_num}...")
            
            # Process each unit
            for unit_num in range(1, 31):
                unit_str = str(unit_num)
                
                # Get current words in app for this unit
                current_words = app_book["units"].get(unit_str, [])
                current_count = len(current_words)
                
                # Get words from document for this unit
                doc_words_list = doc_book["units"].get(unit_str, [])
                doc_count = len(doc_words_list)
                
                # If unit has fewer than 20 words, add placeholders
                if current_count < 20:
                    placeholders_needed = 20 - current_count
                    print(f"  Unit {unit_num}: {current_count} words, adding {placeholders_needed} placeholders")
                    
                    # Add placeholders
                    for i in range(placeholders_needed):
                        placeholder_num = current_count + i + 1
                        placeholder_word = f"placeholder_{unit_num}_{placeholder_num}"
                        
                        current_words.append({
                            "word": placeholder_word,
                            "partOfSpeech": "noun",
                            "meaning": f"Placeholder word {placeholder_num} for Unit {unit_num}",
                            "example": f"This is placeholder {placeholder_num} for learning purposes."
                        })
                    
                    # Update the unit in app data
                    if book_str not in app_data:
                        app_data[book_str] = {"title": f"Book {book_num} - Academic Vocabulary", "units": {}}
                    
                    app_data[book_str]["units"][unit_str] = current_words
                else:
                    print(f"  Unit {unit_num}: {current_count} words ✓")
        
        # Write updated data
        with open('bookData.js', 'w', encoding='utf-8') as f:
            f.write('// Book Data Structure - Extracted from Word Documents\n')
            f.write('// Real vocabulary content from user-provided files\n')
            f.write('// All 4 books with complete vocabulary and placeholders\n\n')
            f.write('const bookData = ')
            f.write(json.dumps(app_data, indent=2, ensure_ascii=False))
            f.write(';\n\n')
            f.write('// Export the data for use in the main script\n')
            f.write('if (typeof module !== \'undefined\' && module.exports) {\n')
            f.write('    module.exports = bookData;\n')
            f.write('}\n')
        
        print("\n✅ bookData.js updated with placeholders!")
        
        # Print final statistics
        total_units = sum(len(book["units"]) for book in app_data.values())
        total_words = sum(
            sum(len(unit) for unit in book["units"].values()) 
            for book in app_data.values()
        )
        print(f"\n📊 Final Statistics:")
        print(f"Total books: {len(app_data)}")
        print(f"Total units: {total_units}")
        print(f"Total words: {total_words}")
        
        for book_num, book_info in app_data.items():
            word_count = sum(len(unit) for unit in book_info["units"].values())
            unit_count = len(book_info["units"])
            print(f"   Book {book_num}: {word_count} words across {unit_count} units")
            
            # Check if all units have 20 words
            incomplete_units = []
            for unit_num, unit_words in book_info["units"].items():
                if len(unit_words) != 20:
                    incomplete_units.append(unit_num)
            
            if incomplete_units:
                print(f"      ⚠️  Incomplete units: {', '.join(incomplete_units)}")
            else:
                print(f"      ✅ All units have 20 words")
        
    except Exception as e:
        print(f"Error updating bookData.js: {e}")

if __name__ == "__main__":
    add_placeholders_to_units()
